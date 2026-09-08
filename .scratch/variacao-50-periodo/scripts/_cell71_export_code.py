# ==============================================================================
# 3.10 EXPORTAÇÃO DOS ARQUIVOS EXCEL MENSAIS
# ==============================================================================
# Garante que a raiz do projeto seja a pasta pai caso o notebook rode dentro da pasta /app
RAIZ_PROJETO = Path.cwd().parent if Path.cwd().name == 'app' else Path.cwd()
PASTA_SAIDA = RAIZ_PROJETO / 'data' / 'outputs' / 'monthly'
PASTA_SAIDA.mkdir(parents=True, exist_ok=True)

DATA_EXPORTACAO = datetime.now().strftime("%Y_%m_%d_%H%M%S")
CAMINHO_ARQUIVO = PASTA_SAIDA / f'{DATA_EXPORTACAO}_indicadores_mensais.xlsx'

# Preparar os dados em português para o Excel
df_exportacao = preparar_dataframe_para_excel(df_exportacao_mensal)

# ------------------------------------------------------------------------------
# Colunas auxiliares "_var_" (ticket 04 de .scratch/variacao-50-periodo/): uma
# por indicador coberto pela regra de variação da 3.8, agrupadas no fim da aba
# "Indicadores Mensais". Ficam ocultas (mas desocultáveis) e alimentam a
# formatação condicional laranja aplicada depois da exportação, mais abaixo
# nesta mesma célula. Acréscimo puro - nenhuma coluna existente muda de nome,
# cabeçalho, ordem ou tipo.
# O cabeçalho exportado do indicador inclui a unidade entre parênteses (ex.:
# "Preço Unitário do Leite (R$/litro)") - vem do mapa técnico -> português da
# 3.9, não do rótulo "nu" da configuração da 3.8. Usar o rótulo nu aqui faria a
# formatação condicional nunca encontrar a coluna e virar um no-op silencioso.
MAPA_TECNICO_PARA_PT_MENSAL = dict(zip(COLUNAS_TECNICAS_MENSAIS_EM_ORDEM, NOMES_PORTUGUES_EM_ORDEM_MENSAIS))

COLUNAS_VAR_MENSAL = {}  # nome da coluna auxiliar -> cabeçalho exportado do indicador correspondente
for _coluna, (_rotulo, _unidade, _formato) in INDICADORES_VARIACAO_MENSAL.items():
    _nome_var = f"_var_{_coluna}"
    df_exportacao[_nome_var] = (
        df_variacao_anotado_mensal[f"variacao_{_coluna}"]
        .reindex(df_exportacao_mensal.index)
        .to_numpy()
    )
    if _coluna not in MAPA_TECNICO_PARA_PT_MENSAL:
        raise KeyError(
            f"Indicador '{_coluna}' não está em COLUNAS_TECNICAS_MENSAIS_EM_ORDEM (3.9) - "
            "a coloração da 3.10 não vai achar a coluna correspondente na aba exportada."
        )
    COLUNAS_VAR_MENSAL[_nome_var] = MAPA_TECNICO_PARA_PT_MENSAL[_coluna]

# Abas novas da regra de variação mensal (tickets 03 e 05): sem preenchimento de
# numéricos vazios com zero, para "Não avaliado" chegar como texto, nunca "0".
df_exportacao_variacoes = preparar_dataframe_para_excel(df_variacoes)
df_exportacao_variacoes_consultor = preparar_dataframe_para_excel(df_variacoes_por_consultor)

# Auditoria da forrageira própria: uma aba por lançamento, uma por lote e uma de
# pendências. As três são montadas na seção 3.4.2.5, que roda antes de
# df_dim_property existir, então o nome da fazenda é colado aqui — sem ele o
# consultor recebe só o UUID e não consegue cobrar o lançamento de ninguém.
COLUNAS_IDENTIFICACAO_FORRAGEIRA = ["id_property", "property_name", "labor_rural_code"]


def identificar_fazenda(df: pd.DataFrame) -> pd.DataFrame:
    """Acrescenta nome e código da fazenda e os traz para as primeiras colunas."""

    if "id_property" not in df.columns:
        return df

    colunas_dim = [
        coluna
        for coluna in COLUNAS_IDENTIFICACAO_FORRAGEIRA
        if coluna in df_dim_property.columns
    ]

    df = df.drop(
        columns=[c for c in colunas_dim if c != "id_property" and c in df.columns],
        errors="ignore",
    ).merge(
        df_dim_property[colunas_dim].drop_duplicates("id_property"),
        on="id_property",
        how="left",
        validate="m:1",
    )

    identificacao = [c for c in colunas_dim if c in df.columns]

    return df[identificacao + [c for c in df.columns if c not in identificacao]]


df_exportacao_forrageira_lancamento = preparar_dataframe_para_excel(identificar_fazenda(df_auditoria_forrageira))
df_exportacao_forrageira_lote = preparar_dataframe_para_excel(identificar_fazenda(df_auditoria_forrageira_lote))
df_exportacao_forrageira_pendencia = preparar_dataframe_para_excel(identificar_fazenda(df_pendencia_custo_forrageira))

# Exportar com a formatação padrão
exportar_varias_abas_xlsx(
    abas={
        'Indicadores Mensais': df_exportacao,
        'Variações': df_exportacao_variacoes,
        'Variações por Consultor': df_exportacao_variacoes_consultor,
        'Forrageira por Lançamento': df_exportacao_forrageira_lancamento,
        'Forrageira por Lote': df_exportacao_forrageira_lote,
        'Pendências de Custo': df_exportacao_forrageira_pendencia,
    },
    caminho_saida=CAMINHO_ARQUIVO,
    fonte='Aptos'
)

# ------------------------------------------------------------------------------
# Coloração laranja pós-exportação (ticket 04 de .scratch/variacao-50-periodo/):
# formatação condicional por fórmula via openpyxl, lendo as colunas "_var_" -
# não estilo célula a célula, porque a aba tem ~168 colunas x ~13 mil linhas.
# Roda DEPOIS de exportar_varias_abas_xlsx porque essa função já recarrega e
# resalva o arquivo (aplicar_estilo_listrado_xlsx, em functions/) - fazer antes
# seria sobrescrito. functions/ não é tocado; esta etapa fica local ao notebook.
# ------------------------------------------------------------------------------
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

_wb_cor = load_workbook(CAMINHO_ARQUIVO)
_ws_indicadores = _wb_cor['Indicadores Mensais']
_preenchimento_laranja = PatternFill(start_color='FFA500', end_color='FFA500', fill_type='solid')

_colunas_por_nome = {celula.value: idx + 1 for idx, celula in enumerate(_ws_indicadores[1])}
_max_linha = _ws_indicadores.max_row
_regras_aplicadas = 0

for _nome_var, _rotulo in COLUNAS_VAR_MENSAL.items():
    if _nome_var not in _colunas_por_nome or _rotulo not in _colunas_por_nome:
        raise KeyError(
            f"Coloração da 3.10: coluna '{_nome_var}' ou '{_rotulo}' não encontrada no "
            f"cabeçalho exportado. Colunas disponíveis (amostra): "
            f"{list(_colunas_por_nome)[:5]} ... {list(_colunas_por_nome)[-10:]}"
        )
    _letra_var = get_column_letter(_colunas_por_nome[_nome_var])
    _letra_indicador = get_column_letter(_colunas_por_nome[_rotulo])

    # Coluna auxiliar oculta por padrão - o auditor consegue desocultá-la.
    _ws_indicadores.column_dimensions[_letra_var].hidden = True

    if _max_linha >= 2:
        _intervalo = f"{_letra_indicador}2:{_letra_indicador}{_max_linha}"
        _formula = f'OR(${_letra_var}2="Aumento",${_letra_var}2="Redução")'
        _ws_indicadores.conditional_formatting.add(
            _intervalo,
            FormulaRule(formula=[_formula], fill=_preenchimento_laranja),
        )
        _regras_aplicadas += 1

if _regras_aplicadas != len(COLUNAS_VAR_MENSAL):
    raise AssertionError(
        f"Coloração da 3.10: só {_regras_aplicadas} de {len(COLUNAS_VAR_MENSAL)} regras "
        "condicionais foram aplicadas."
    )

_wb_cor.save(CAMINHO_ARQUIVO)
print(f"Coloração aplicada: {_regras_aplicadas} colunas de indicador com formatação condicional laranja.")

print('Exportação mensal concluída com sucesso.')
print(f'Linhas exportadas: {len(df_exportacao):,}')
print(f'Aba Variações: {len(df_exportacao_variacoes):,} linhas | '
      f'Variações por Consultor: {len(df_exportacao_variacoes_consultor):,} linhas')
print(f'Arquivo: {CAMINHO_ARQUIVO}')
